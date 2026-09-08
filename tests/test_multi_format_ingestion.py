"""
Multi-Format Ingestion Test Suite
=================================
Validates multi-format detection, extraction, chunking, Lakehouse storage,
structural entity/relationship generation, and Apache AGE ingestion across
PDF, TXT, MD, DOCX, HTML, XML, CSV, TSV, XLSX, JSON, JSONL, Parquet, Feather, YAML, SQL, RTF.
"""

import sys
import os
import io
import json
import unittest
import tempfile
from pathlib import Path

# Add project root to sys.path
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from backend.storage.lakehouse import LocalLakehouseStorageService
from backend.storage.models import ProcessingStatus
from backend.extraction.file_detector import FileDetector, SupportedFormat, FileDetectionError
from backend.extraction.document_extractor import DocumentExtractor, DocumentExtractionError
from backend.extraction.chunker import DocumentChunker
from backend.context.extractor import ContextExtractor
from backend.pipeline import PDFContextPipeline
from backend.graph.age_client import AGEClient
from backend.graph.graph_service import GraphService
from fastapi.testclient import TestClient
from backend.main import app

class TestMultiFormatIngestion(unittest.TestCase):
    def setUp(self):
        self.test_dir = Path(__file__).parent.parent / "backend" / "test_multi_format_storage"
        self.storage = LocalLakehouseStorageService(root_dir=self.test_dir)
        self.extractor = DocumentExtractor()
        self.chunker = DocumentChunker()
        self.context_extractor = ContextExtractor()
        self.graph_service = GraphService(graph_name="test_multi_format_kg")
        self.pipeline = PDFContextPipeline(
            storage=self.storage,
            extractor=self.extractor,
            chunker=self.chunker,
            context_extractor=self.context_extractor,
            graph_service=self.graph_service
        )
        self.api_client = TestClient(app)

    def tearDown(self):
        # Clean up test storage directory
        import shutil
        if self.test_dir.exists():
            shutil.rmtree(self.test_dir)

    # -------------------------------------------------------------------------
    # 1. File Type Detection Tests
    # -------------------------------------------------------------------------
    def test_01_file_detector_supported_formats(self):
        test_cases = [
            ("report.pdf", SupportedFormat.PDF, "application/pdf"),
            ("notes.txt", SupportedFormat.TXT, "text/plain"),
            ("README.md", SupportedFormat.MARKDOWN, "text/markdown"),
            ("document.docx", SupportedFormat.DOCX, "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
            ("page.html", SupportedFormat.HTML, "text/html"),
            ("page.htm", SupportedFormat.HTML, "text/html"),
            ("data.xml", SupportedFormat.XML, "application/xml"),
            ("customers.csv", SupportedFormat.CSV, "text/csv"),
            ("metrics.tsv", SupportedFormat.TSV, "text/tab-separated-values"),
            ("payload.json", SupportedFormat.JSON, "application/json"),
            ("events.jsonl", SupportedFormat.JSONL, "application/x-ndjson"),
            ("events.ndjson", SupportedFormat.JSONL, "application/x-ndjson"),
            ("finance.xlsx", SupportedFormat.XLSX, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
            ("legacy.xls", SupportedFormat.XLS, "application/vnd.ms-excel"),
            ("analytics.parquet", SupportedFormat.PARQUET, "application/vnd.apache.parquet"),
            ("data.feather", SupportedFormat.FEATHER, "application/vnd.apache.arrow.feather"),
            ("config.yaml", SupportedFormat.YAML, "application/x-yaml"),
            ("config.yml", SupportedFormat.YAML, "application/x-yaml"),
            ("schema.sql", SupportedFormat.SQL, "application/sql"),
            ("letter.rtf", SupportedFormat.RTF, "application/rtf"),
        ]
        for filename, expected_fmt, expected_mime in test_cases:
            fmt, mime, cat = FileDetector.detect_format(filename)
            self.assertEqual(fmt, expected_fmt, f"Failed for {filename}")
            self.assertEqual(mime, expected_mime, f"Failed MIME for {filename}")

    def test_02_file_detector_unsupported_rejection(self):
        unsupported = ["virus.exe", "image.png", "music.mp3", "archive.tar.gz", "script.sh"]
        for fn in unsupported:
            with self.assertRaises(FileDetectionError):
                FileDetector.detect_format(fn)

    # -------------------------------------------------------------------------
    # 2. Format-Specific Extraction Tests
    # -------------------------------------------------------------------------
    def test_03_txt_extraction(self):
        txt_content = b"Context Lakehouse Architecture\nApache AGE extends PostgreSQL with Graph capabilities.\nPython is used for data engineering."
        doc = self.extractor.extract_from_bytes(txt_content, "doc_txt_1", "architecture.txt")
        self.assertEqual(doc.total_pages, 1)
        self.assertGreater(doc.total_characters, 50)
        self.assertIn("Context Lakehouse", doc.pages[0].text)

    def test_04_markdown_extraction(self):
        md_content = b"# Apache AGE Guide\n\n## Overview\nApache AGE is a graph database extension for PostgreSQL.\n\n- Entity extraction\n- openCypher support"
        doc = self.extractor.extract_from_bytes(md_content, "doc_md_1", "guide.md")
        self.assertEqual(doc.total_pages, 1)
        self.assertIn("Apache AGE", doc.pages[0].text)
        self.assertEqual(doc.metadata.get("format"), "markdown")

    def test_05_csv_extraction(self):
        csv_content = b"customer_id,name,city,total_spend\n1,Ravi,Hyderabad,1500.50\n2,Priya,Bengaluru,2400.00\n3,Amit,Mumbai,850.00\n"
        doc = self.extractor.extract_from_bytes(csv_content, "doc_csv_1", "customers.csv")
        self.assertEqual(doc.metadata.get("record_count"), 3)
        self.assertEqual(doc.metadata.get("column_count"), 4)
        self.assertIn("customer_id", doc.metadata.get("columns"))
        self.assertEqual(doc.metadata.get("column_types").get("customer_id"), "Integer")
        self.assertEqual(doc.metadata.get("column_types").get("total_spend"), "Float")

    def test_06_tsv_extraction(self):
        tsv_content = b"metric_id\tname\tcategory\n101\tchurn_rate\tretention\n102\tmrr\trevenue\n"
        doc = self.extractor.extract_from_bytes(tsv_content, "doc_tsv_1", "metrics.tsv")
        self.assertEqual(doc.metadata.get("record_count"), 2)
        self.assertEqual(doc.metadata.get("column_count"), 3)
        self.assertIn("metric_id", doc.metadata.get("columns"))

    def test_07_json_extraction(self):
        json_obj = {
            "application": "Context Lakehouse",
            "version": "2.0.0",
            "database": "PostgreSQL",
            "extension": "Apache AGE",
            "features": ["AST Parsing", "Provenance Tracking", "Multi-Format Ingestion"]
        }
        json_bytes = json.dumps(json_obj).encode("utf-8")
        doc = self.extractor.extract_from_bytes(json_bytes, "doc_json_1", "app_config.json")
        self.assertIn("application", doc.metadata.get("schema_keys"))
        self.assertIn("Context Lakehouse", doc.pages[0].text)

    def test_08_jsonl_extraction(self):
        jsonl_lines = [
            json.dumps({"event_id": 1, "user": "Ravi", "action": "login"}),
            json.dumps({"event_id": 2, "user": "Priya", "action": "upload_file"}),
            json.dumps({"event_id": 3, "user": "Amit", "action": "export_graph"})
        ]
        jsonl_bytes = "\n".join(jsonl_lines).encode("utf-8")
        doc = self.extractor.extract_from_bytes(jsonl_bytes, "doc_jsonl_1", "events.jsonl")
        self.assertEqual(doc.metadata.get("record_count"), 3)
        self.assertIn("event_id", doc.metadata.get("schema_keys"))

    def test_09_xml_extraction(self):
        xml_content = b"""<?xml version="1.0" encoding="UTF-8"?>
<catalog>
    <book id="bk101">
        <author>Rowthu Kalyani</author>
        <title>Context Lakehouse Architecture</title>
        <technology>Apache AGE</technology>
    </book>
</catalog>"""
        doc = self.extractor.extract_from_bytes(xml_content, "doc_xml_1", "catalog.xml")
        self.assertEqual(doc.metadata.get("root_tag"), "catalog")
        self.assertGreater(doc.metadata.get("element_count"), 3)
        self.assertIn("Apache AGE", doc.pages[0].text)

    def test_10_html_extraction(self):
        html_content = b"""<!DOCTYPE html>
<html>
<head><title>Context Lakehouse Project</title></head>
<body>
<h1>Context Lakehouse & Apache AGE</h1>
<p>This project builds a unified knowledge graph from structured data and documents.</p>
</body>
</html>"""
        doc = self.extractor.extract_from_bytes(html_content, "doc_html_1", "index.html")
        self.assertIn("Context Lakehouse", doc.pages[0].text)
        self.assertIn("Apache AGE", doc.pages[0].text)

    def test_11_docx_extraction(self):
        # Create minimal in-memory docx using python-docx
        import docx
        doc_obj = docx.Document()
        doc_obj.add_heading("Context Lakehouse Report", level=1)
        doc_obj.add_paragraph("Apache AGE is a graph extension for PostgreSQL.")
        table = doc_obj.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Tech"
        table.cell(0, 1).text = "Role"
        table.cell(1, 0).text = "DuckDB"
        table.cell(1, 1).text = "Analytical Query Engine"
        
        docx_buf = io.BytesIO()
        doc_obj.save(docx_buf)
        docx_bytes = docx_buf.getvalue()

        doc = self.extractor.extract_from_bytes(docx_bytes, "doc_docx_1", "report.docx")
        self.assertIn("Context Lakehouse", doc.pages[0].text)
        self.assertIn("Apache AGE", doc.pages[0].text)
        self.assertIn("DuckDB", doc.pages[0].text)

    def test_12_xlsx_extraction(self):
        # Create in-memory xlsx using openpyxl
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Technologies"
        ws.append(["Name", "Category", "Engine"])
        ws.append(["Apache AGE", "Graph DB", "PostgreSQL"])
        ws.append(["DuckDB", "OLAP", "C++"])

        xlsx_buf = io.BytesIO()
        wb.save(xlsx_buf)
        xlsx_bytes = xlsx_buf.getvalue()

        doc = self.extractor.extract_from_bytes(xlsx_bytes, "doc_xlsx_1", "tech_stack.xlsx")
        self.assertEqual(doc.metadata.get("sheet_count"), 1)
        self.assertEqual(doc.metadata.get("record_count"), 2)
        self.assertIn("Apache AGE", doc.pages[0].text)

    def test_13_parquet_and_feather_extraction(self):
        import pyarrow as pa
        import pyarrow.parquet as pq
        import pyarrow.feather as feather

        table = pa.Table.from_arrays(
            [pa.array([1, 2, 3]), pa.array(["PostgreSQL", "Apache AGE", "FastAPI"]), pa.array([10.5, 20.0, 30.5])],
            names=["id", "technology", "score"]
        )

        # Test Parquet
        pq_buf = io.BytesIO()
        pq.write_table(table, pq_buf)
        pq_bytes = pq_buf.getvalue()

        doc_pq = self.extractor.extract_from_bytes(pq_bytes, "doc_pq_1", "data.parquet")
        self.assertEqual(doc_pq.metadata.get("record_count"), 3)
        self.assertIn("technology", doc_pq.metadata.get("columns"))
        self.assertIn("Apache AGE", doc_pq.pages[0].text)

        # Test Feather
        fea_buf = io.BytesIO()
        feather.write_feather(table, fea_buf)
        fea_bytes = fea_buf.getvalue()

        doc_fea = self.extractor.extract_from_bytes(fea_bytes, "doc_fea_1", "data.feather")
        self.assertEqual(doc_fea.metadata.get("record_count"), 3)
        self.assertIn("technology", doc_fea.metadata.get("columns"))

    def test_14_yaml_extraction(self):
        yaml_content = b"""
pipeline:
  name: Context Lakehouse
  version: 2.0
  components:
    - Apache AGE
    - PostgreSQL
    - Streamlit
"""
        doc = self.extractor.extract_from_bytes(yaml_content, "doc_yaml_1", "pipeline.yaml")
        self.assertIn("pipeline", doc.metadata.get("top_keys"))
        self.assertIn("Apache AGE", doc.pages[0].text)

    def test_15_sql_extraction(self):
        sql_content = b"""
CREATE TABLE customers (
    customer_id INT PRIMARY KEY,
    name VARCHAR(100),
    city VARCHAR(100)
);

SELECT * FROM customers WHERE city = 'Hyderabad';
"""
        doc = self.extractor.extract_from_bytes(sql_content, "doc_sql_1", "queries.sql")
        self.assertEqual(doc.metadata.get("sql_statement_count"), 2)
        self.assertIn("customers", doc.pages[0].text)

    def test_16_rtf_extraction(self):
        rtf_content = b"{\\rtf1\\ansi\\deff0 {\\fonttbl {\\f0 Courier;}}\\f0\\fs24 Context Lakehouse with Apache AGE and PostgreSQL.}"
        doc = self.extractor.extract_from_bytes(rtf_content, "doc_rtf_1", "document.rtf")
        self.assertIn("Context Lakehouse", doc.pages[0].text)
        self.assertIn("Apache AGE", doc.pages[0].text)

    # -------------------------------------------------------------------------
    # 3. End-to-End Pipeline & Apache AGE Ingestion Tests
    # -------------------------------------------------------------------------
    def test_17_e2e_csv_pipeline_and_graph_ingestion(self):
        csv_bytes = b"service,database,language\nContext Lakehouse,PostgreSQL,Python\nApache AGE,PostgreSQL,C\nFastAPI,None,Python\n"
        meta = self.pipeline.upload_document(filename="services.csv", file_bytes=csv_bytes)
        self.assertEqual(meta.status, ProcessingStatus.UPLOADED)
        self.assertEqual(meta.file_type, "text/csv")

        # Execute processing pipeline
        res = self.pipeline.process_document(meta.document_id)
        self.assertEqual(res["status"], ProcessingStatus.COMPLETED.value)
        self.assertGreater(res["total_entities"], 0)
        self.assertGreater(res["total_relationships"], 0)

        # Verify Apache AGE Knowledge Graph contains the entities
        cypher_res = self.graph_service.age_client.execute_cypher(
            "MATCH (n) WHERE n.canonical_name = 'Services' OR n.entity_type = 'Column' RETURN count(n) AS cnt",
            columns=["cnt"],
            graph_name=self.graph_service.graph_name
        )
        self.assertGreater(cypher_res[0]["cnt"], 0)

    def test_18_e2e_json_pipeline_and_graph_ingestion(self):
        json_bytes = json.dumps({
            "project": "Context Lakehouse",
            "technologies": ["Apache AGE", "PostgreSQL", "DuckDB", "FastAPI"],
            "status": "Active"
        }).encode("utf-8")

        meta = self.pipeline.upload_document(filename="project_info.json", file_bytes=json_bytes)
        res = self.pipeline.process_document(meta.document_id)
        self.assertEqual(res["status"], ProcessingStatus.COMPLETED.value)
        self.assertGreater(res["total_entities"], 0)

    def test_19_idempotent_duplicate_upload(self):
        content = b"Deterministic duplicate content for test."
        m1 = self.pipeline.upload_document(filename="file_a.txt", file_bytes=content)
        m2 = self.pipeline.upload_document(filename="file_b.txt", file_bytes=content)
        self.assertEqual(m1.document_id, m2.document_id)

    # -------------------------------------------------------------------------
    # 4. Error Handling & Edge Cases
    # -------------------------------------------------------------------------
    def test_20_corrupted_json_handling(self):
        bad_json = b"{'invalid_json': missing_quotes}"
        with self.assertRaises(DocumentExtractionError):
            self.extractor.extract_from_bytes(bad_json, "doc_bad", "bad.json")

    def test_21_corrupted_xml_handling(self):
        bad_xml = b"<root><unclosed></root>"
        with self.assertRaises(DocumentExtractionError):
            self.extractor.extract_from_bytes(bad_xml, "doc_bad", "bad.xml")

    def test_22_empty_file_handling(self):
        with self.assertRaises(DocumentExtractionError):
            self.extractor.extract_from_bytes(b"", "doc_empty", "empty.txt")

    # -------------------------------------------------------------------------
    # 5. REST API Multi-Format Upload Endpoint Tests
    # -------------------------------------------------------------------------
    def test_23_api_upload_multi_format_success(self):
        # Test CSV upload
        csv_file = io.BytesIO(b"id,name\n1,Test")
        res = self.api_client.post(
            "/documents/upload",
            files={"file": ("api_test.csv", csv_file, "text/csv")}
        )
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.json()["file_type"], "text/csv")

        # Test JSON upload
        json_file = io.BytesIO(b'{"key": "value"}')
        res = self.api_client.post(
            "/documents/upload",
            files={"file": ("api_test.json", json_file, "application/json")}
        )
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.json()["file_type"], "application/json")

    def test_24_api_upload_unsupported_format_rejected(self):
        exe_file = io.BytesIO(b"MZ binary executable content")
        res = self.api_client.post(
            "/documents/upload",
            files={"file": ("malware.exe", exe_file, "application/octet-stream")}
        )
        self.assertEqual(res.status_code, 400)
        self.assertIn("Unsupported file format", res.json()["detail"])

    def test_25_excel_e2e_kg_no_orphan_concept_nodes(self):
        """Test Excel ingestion creates connected structural graph without orphan Concept nodes."""
        import pandas as pd
        df = pd.DataFrame({
            "ID": [101, 102],
            "Name": ["Ravi", "Anu"],
            "Department": ["Engineering", "HR"]
        })
        buf = io.BytesIO()
        df.to_excel(buf, index=False, sheet_name="Sheet1")
        excel_bytes = buf.getvalue()

        meta = self.pipeline.upload_document(filename="Book1.xlsx", file_bytes=excel_bytes)
        res = self.pipeline.process_document(meta.document_id)
        self.assertEqual(res["status"], ProcessingStatus.COMPLETED.value)
        self.assertEqual(res["total_entities"], 5)
        self.assertEqual(res["total_relationships"], 4)

        # Check in storage context tier
        ctx = self.storage.get_context(meta.document_id)
        self.assertIsNotNone(ctx)
        entities = ctx.get("entities", [])
        self.assertEqual(len(entities), 5)
        cnames = {e.get("canonical_name") for e in entities}
        self.assertIn("Book1.xlsx", cnames)
        self.assertIn("Book1.xlsx::Sheet1", cnames)
        self.assertIn("Book1.xlsx::Sheet1.ID", cnames)
        self.assertIn("Book1.xlsx::Sheet1.Name", cnames)
        self.assertIn("Book1.xlsx::Sheet1.Department", cnames)

        # Confirm NO orphan concept nodes
        orphan_concepts = [e.get("canonical_name") for e in entities if e.get("type") == "Concept"]
        self.assertEqual(orphan_concepts, [])

if __name__ == "__main__":
    unittest.main()
